import os
import json
import re
from pathlib import Path
from datetime import datetime
from werkzeug.utils import secure_filename
from config import Config
from database import get_db

def extract_text_from_file(file_path: str) -> str:
    """Extracts text content safely from PDF, DOCX, DOC, XLSX, or TXT files."""
    if not file_path:
        return ""
    full_path = Path(file_path)
    if not full_path.is_absolute():
        full_path = Config.BASE_DIR / file_path
        
    if not full_path.exists():
        return f"[File not found: {file_path}]"

    suffix = full_path.suffix.lower()
    
    # 1. Plain Text / Markdown / Code
    if suffix in ['.txt', '.md', '.json', '.py', '.html', '.css', '.csv']:
        try:
            with open(full_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            return f"[Error reading text file: {e}]"
            
    # 2. PDF Files (PyMuPDF / fitz with pypdf fallback)
    elif suffix == '.pdf':
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(full_path)
            text_chunks = []
            for page in doc:
                text_chunks.append(page.get_text())
            doc.close()
            return "\n".join(text_chunks)
        except Exception:
            try:
                import pypdf
                reader = pypdf.PdfReader(full_path)
                return "\n".join([page.extract_text() or '' for page in reader.pages])
            except Exception as e:
                return f"[Error extracting PDF text: {e}]"
            
    # 3. Modern Word Documents (.docx)
    elif suffix == '.docx':
        try:
            import docx
            doc = docx.Document(full_path)
            text_chunks = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_chunks.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_chunks.append(" | ".join(row_cells))
            return "\n".join(text_chunks)
        except Exception as e:
            return f"[Error extracting DOCX text: {e}]"

    # 4. Legacy Word Documents (.doc) & RTF
    elif suffix in ['.doc', '.rtf']:
        try:
            with open(full_path, 'rb') as f:
                content = f.read()
            strings = re.findall(rb'[\x20-\x7E\t\n\r]{4,}', content)
            decoded = [s.decode('utf-8', errors='ignore').strip() for s in strings if len(s.decode('utf-8', errors='ignore').strip()) > 3]
            clean_lines = [line for line in decoded if not line.startswith(('CompObj', 'WordDocument', 'SummaryInformation', 'DocumentSummaryInformation', '{\\rtf'))]
            if clean_lines:
                return "\n".join(clean_lines)
            return f"[Legacy Word Doc binary content: {len(decoded)} tokens]"
        except Exception as e:
            return f"[Error extracting legacy Word DOC text: {e}]"

    # 5. XLSX / Excel Spreadsheet Files
    elif suffix in ['.xlsx', '.xls']:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(full_path, data_only=True)
            lines = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                lines.append(f"--- Sheet: {sheet} ---")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            return "\n".join(lines)
        except Exception as e:
            return f"[Error extracting Excel data: {e}]"
            
    return f"[Unsupported file format: {suffix}]"


def init_documents_table():
    """Ensure the documents table exists in database with all comparative verification columns."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS documents (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        file_path TEXT NOT NULL,
        file_name TEXT NOT NULL,
        file_type TEXT NOT NULL,
        file_size_kb REAL DEFAULT 0.0,
        category TEXT NOT NULL,
        project_id INTEGER,
        uploaded_by INTEGER NOT NULL,
        submission_id INTEGER,
        intern_id INTEGER,
        intern_name TEXT,
        sample_file_path TEXT,
        sample_file_name TEXT,
        extracted_text_preview TEXT,
        extracted_sections_json TEXT,
        formatting_checks_json TEXT,
        missing_sections_json TEXT,
        comparison_details_json TEXT,
        format_comparison_matrix_json TEXT,
        proper_format_template TEXT,
        intern_actionable_feedback TEXT,
        format_score INTEGER DEFAULT 0,
        format_status TEXT DEFAULT 'Processing',
        approval_status TEXT DEFAULT 'Pending Review',
        executive_verdict TEXT,
        approved_by INTEGER,
        approved_at TIMESTAMP,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE SET NULL,
        FOREIGN KEY (uploaded_by) REFERENCES users(id) ON DELETE CASCADE,
        FOREIGN KEY (submission_id) REFERENCES submissions(id) ON DELETE SET NULL,
        FOREIGN KEY (approved_by) REFERENCES users(id) ON DELETE SET NULL
    );
    ''')
    for col, col_type in [
        ("submission_id", "INTEGER"),
        ("intern_id", "INTEGER"),
        ("intern_name", "TEXT"),
        ("sample_file_path", "TEXT"),
        ("sample_file_name", "TEXT"),
        ("missing_sections_json", "TEXT"),
        ("comparison_details_json", "TEXT"),
        ("format_comparison_matrix_json", "TEXT"),
        ("proper_format_template", "TEXT"),
        ("intern_actionable_feedback", "TEXT"),
    ]:
        try:
            cursor.execute(f"ALTER TABLE documents ADD COLUMN {col} {col_type};")
            conn.commit()
        except Exception:
            pass
    conn.commit()
    conn.close()


def save_and_verify_document(
    intern_file_obj=None,
    sample_file_obj=None,
    submission_id: int = None,
    title: str = "",
    category: str = "Milestone Report",
    project_id: int = None,
    uploaded_by: int = 1
) -> tuple:
    """
    Ingests an intern's submitted document AND the mentor's uploaded sample reference document,
    executes AI comparative structural verification, detects missing sections and formatting omissions,
    and stores full audit findings.
    """
    from gemini_service import compare_document_against_sample_ai, validate_document_structure_ai
    
    init_documents_table()
    
    docs_folder = Config.BASE_DIR / 'uploads' / 'documents'
    samples_folder = Config.BASE_DIR / 'uploads' / 'samples'
    docs_folder.mkdir(parents=True, exist_ok=True)
    samples_folder.mkdir(parents=True, exist_ok=True)
    
    conn = get_db()
    cursor = conn.cursor()
    
    intern_text = ""
    intern_file_path = ""
    intern_file_name = ""
    intern_file_type = "docx"
    intern_file_size_kb = 0.0
    intern_id = None
    intern_name = "Intern"
    proj_title = "General Platform Project"
    
    # 1. Case A: Selected from existing intern submission
    if submission_id:
        cursor.execute('''
        SELECT s.*, u.id as intern_user_id, u.full_name as intern_user_name,
               p.id as proj_id, p.title as proj_title
        FROM submissions s
        JOIN assignments a ON s.assignment_id = a.id
        JOIN projects p ON a.project_id = p.id
        JOIN users u ON s.intern_id = u.id
        WHERE s.id = ?
        ''', (submission_id,))
        sub_row = cursor.fetchone()
        if sub_row:
            sub = dict(sub_row)
            intern_id = sub['intern_user_id']
            intern_name = sub['intern_user_name']
            project_id = sub['proj_id']
            proj_title = sub['proj_title']
            intern_file_path = sub['file_path']
            intern_file_name = sub['title']
            intern_file_type = sub.get('file_type', 'docx')
            intern_text = extract_text_from_file(intern_file_path)
            
            full_path = Config.BASE_DIR / intern_file_path
            if full_path.exists():
                intern_file_size_kb = round(os.path.getsize(str(full_path)) / 1024.0, 1)
            else:
                intern_file_size_kb = 120.0

    # 1. Case B: Intern file uploaded directly
    if intern_file_obj and intern_file_obj.filename:
        original_filename = secure_filename(intern_file_obj.filename)
        intern_file_type = original_filename.rsplit('.', 1)[1].lower() if '.' in original_filename else 'docx'
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        saved_filename = f"intern_doc_{uploaded_by}_{timestamp}_{original_filename}"
        full_save_path = docs_folder / saved_filename
        intern_file_obj.save(str(full_save_path))
        
        intern_file_path = f"uploads/documents/{saved_filename}"
        intern_file_name = original_filename
        intern_file_size_kb = round(os.path.getsize(str(full_save_path)) / 1024.0, 1)
        intern_text = extract_text_from_file(intern_file_path)
        if not title:
            title = original_filename.rsplit('.', 1)[0].replace('_', ' ')

    # 2. Mentor Sample Document
    sample_text = ""
    sample_file_path = ""
    sample_file_name = ""
    
    if sample_file_obj and sample_file_obj.filename:
        orig_sample_name = secure_filename(sample_file_obj.filename)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        saved_sample_name = f"sample_bench_{uploaded_by}_{timestamp}_{orig_sample_name}"
        sample_save_path = samples_folder / saved_sample_name
        sample_file_obj.save(str(sample_save_path))
        
        sample_file_path = f"uploads/samples/{saved_sample_name}"
        sample_file_name = orig_sample_name
        sample_text = extract_text_from_file(sample_file_path)
    else:
        # Check if project has an existing reference standard in sample_documents or documents
        if project_id:
            cursor.execute("SELECT * FROM sample_documents WHERE project_id = ? ORDER BY id DESC LIMIT 1", (project_id,))
            s_row = cursor.fetchone()
            if s_row:
                s_dict = dict(s_row)
                sample_file_path = s_dict.get('file_path', '')
                sample_file_name = s_dict.get('title', 'Reference Standard')
                sample_text = extract_text_from_file(sample_file_path)
            else:
                cursor.execute("SELECT * FROM documents WHERE project_id = ? AND approval_status = 'Approved' ORDER BY id DESC LIMIT 1", (project_id,))
                d_row = cursor.fetchone()
                if d_row:
                    d_dict = dict(d_row)
                    sample_file_path = d_dict.get('file_path', '')
                    sample_file_name = d_dict.get('title', 'Project Approved Standard')
                    sample_text = extract_text_from_file(sample_file_path)

    if not sample_text:
        # Generate sample standard text structure
        sample_file_name = "Enterprise Milestone Report Standard.docx"
        sample_text = f"""# Master Engineering Standard & Benchmark: {proj_title}
## 1. Executive Summary & Abstract
Comprehensive summary of project background, objectives, and scope.
## 2. System Architecture & Component Diagram
High-level modular architecture diagram, technology stack, and component interactions.
## 3. Data Pipeline & Schema Design
Database schemas, indexing strategy, and data transformation pipeline.
## 4. Technical Methodology & Algorithms
Core algorithm formulations, mathematical models, and workflow graphs.
## 5. Implementation Milestones (Milestones 1 through 6)
Milestone 1: Environment & Baseline Setup
Milestone 2: Data Ingestion & Preprocessing
Milestone 3: Core Model Training & Agent Logic
Milestone 4: API Integration & Backend Endpoints
Milestone 5: Frontend UI & Interactive Testing
Milestone 6: Verification, Containerization & CI/CD
## 6. Testing, Performance & Latency Benchmarks
Unit test coverage, load testing, security validation, and pass rates.
## 7. References, Appendices & Citations
Authoritative literature and technical standards."""

    if project_id and not proj_title:
        cursor.execute("SELECT title FROM projects WHERE id = ?", (project_id,))
        p_row = cursor.fetchone()
        if p_row:
            proj_title = p_row['title']

    # 3. Run Comparative Structural Inspection AI Agent
    analysis = compare_document_against_sample_ai(
        intern_doc_text=intern_text,
        intern_filename=intern_file_name or "Intern_Deliverable.docx",
        sample_doc_text=sample_text,
        sample_filename=sample_file_name,
        project_title=proj_title,
        doc_category=category
    )
    
    format_score = analysis.get("format_score", 45)
    format_status = analysis.get("format_status", "Needs Revision")
    rec = analysis.get("approval_recommendation", "Needs Formatting Revisions")
    
    if "Ready" in rec or format_score >= 80:
        approval_status = "Ready for Approval"
    elif "Revision" in rec or format_score >= 40:
        approval_status = "Needs Revision"
    else:
        approval_status = "Pending Review"
        
    doc_title = title.strip() or analysis.get("document_title") or intern_file_name.rsplit('.', 1)[0]
    
    # 4. Insert into documents table
    cursor.execute('''
    INSERT INTO documents (
        title, file_path, file_name, file_type, file_size_kb, category,
        project_id, uploaded_by, submission_id, intern_id, intern_name,
        sample_file_path, sample_file_name, extracted_text_preview,
        extracted_sections_json, formatting_checks_json, missing_sections_json,
        comparison_details_json, format_comparison_matrix_json, proper_format_template,
        intern_actionable_feedback, format_score, format_status, approval_status,
        executive_verdict
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        doc_title,
        intern_file_path,
        intern_file_name or doc_title,
        intern_file_type,
        intern_file_size_kb,
        category,
        project_id if project_id else None,
        uploaded_by,
        submission_id,
        intern_id,
        intern_name,
        sample_file_path,
        sample_file_name,
        intern_text[:600],
        json.dumps(analysis.get("section_comparisons", [])),
        json.dumps(analysis.get("formatting_checks", [])),
        json.dumps(analysis.get("missing_sections", [])),
        json.dumps(analysis.get("section_comparisons", [])),
        json.dumps(analysis.get("format_comparison_matrix", [])),
        analysis.get("proper_format_template", ""),
        json.dumps(analysis.get("intern_actionable_feedback", [])),
        format_score,
        format_status,
        approval_status,
        analysis.get("executive_verdict", "Comparative inspection completed.")
    ))
    doc_id = cursor.lastrowid
    
    # If linked to a submission, update submission status
    if submission_id:
        cursor.execute("UPDATE submissions SET status = ? WHERE id = ?", (
            'Evaluated' if approval_status == 'Ready for Approval' else 'Needs Revision',
            submission_id
        ))
        
    conn.commit()
    conn.close()
    
    return doc_id, analysis


def get_all_documents(project_id: int = None, category: str = None) -> list:
    """Fetch all uploaded documents with author, project name, intern details, and parsed sections."""
    init_documents_table()
    conn = get_db()
    cursor = conn.cursor()
    
    query = '''
    SELECT d.*, u.full_name as uploader_name, p.title as project_title,
           ap.full_name as approver_name
    FROM documents d
    JOIN users u ON d.uploaded_by = u.id
    LEFT JOIN projects p ON d.project_id = p.id
    LEFT JOIN users ap ON d.approved_by = ap.id
    '''
    params = []
    conditions = []
    
    if project_id:
        conditions.append("d.project_id = ?")
        params.append(project_id)
    if category and category != 'All':
        conditions.append("d.category = ?")
        params.append(category)
        
    if conditions:
        query += " WHERE " + " AND ".join(conditions)
        
    query += " ORDER BY d.id DESC"
    
    cursor.execute(query, params)
    raw_docs = [dict(r) for r in cursor.fetchall()]
    conn.close()
    
    # Parse JSON fields
    for doc in raw_docs:
        try:
            doc['extracted_sections'] = json.loads(doc.get('extracted_sections_json') or '[]')
        except Exception:
            doc['extracted_sections'] = []
            
        try:
            doc['formatting_checks'] = json.loads(doc.get('formatting_checks_json') or '[]')
        except Exception:
            doc['formatting_checks'] = []
            
        try:
            doc['missing_sections'] = json.loads(doc.get('missing_sections_json') or '[]')
        except Exception:
            doc['missing_sections'] = []

        try:
            doc['section_comparisons'] = json.loads(doc.get('comparison_details_json') or '[]')
        except Exception:
            doc['section_comparisons'] = []

        try:
            doc['format_comparison_matrix'] = json.loads(doc.get('format_comparison_matrix_json') or '[]')
        except Exception:
            doc['format_comparison_matrix'] = []

        try:
            doc['intern_actionable_feedback'] = json.loads(doc.get('intern_actionable_feedback') or '[]')
        except Exception:
            doc['intern_actionable_feedback'] = []
            
        doc['sections_count'] = len([s for s in doc['extracted_sections'] if s.get('status') == 'Present' or s.get('intern_status') == 'Present'])
        doc['missing_count'] = len(doc.get('missing_sections', []))
        
    return raw_docs


def get_document_by_id(doc_id: int) -> dict:
    """Fetch single document record by ID with full comparative inspection fields."""
    init_documents_table()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    SELECT d.*, u.full_name as uploader_name, p.title as project_title,
           ap.full_name as approver_name
    FROM documents d
    JOIN users u ON d.uploaded_by = u.id
    LEFT JOIN projects p ON d.project_id = p.id
    LEFT JOIN users ap ON d.approved_by = ap.id
    WHERE d.id = ?
    ''', (doc_id,))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        return None
        
    doc = dict(row)
    try:
        doc['extracted_sections'] = json.loads(doc.get('extracted_sections_json') or '[]')
    except Exception:
        doc['extracted_sections'] = []
        
    try:
        doc['formatting_checks'] = json.loads(doc.get('formatting_checks_json') or '[]')
    except Exception:
        doc['formatting_checks'] = []
        
    try:
        doc['missing_sections'] = json.loads(doc.get('missing_sections_json') or '[]')
    except Exception:
        doc['missing_sections'] = []

    try:
        doc['section_comparisons'] = json.loads(doc.get('comparison_details_json') or '[]')
    except Exception:
        doc['section_comparisons'] = []

    try:
        doc['format_comparison_matrix'] = json.loads(doc.get('format_comparison_matrix_json') or '[]')
    except Exception:
        doc['format_comparison_matrix'] = []

    try:
        doc['intern_actionable_feedback'] = json.loads(doc.get('intern_actionable_feedback') or '[]')
    except Exception:
        doc['intern_actionable_feedback'] = []

    doc['sections_count'] = len([s for s in doc['extracted_sections'] if s.get('status') == 'Present' or s.get('intern_status') == 'Present'])
    doc['missing_count'] = len(doc.get('missing_sections', []))
        
    return doc


def approve_document(doc_id: int, approved_by_user_id: int) -> bool:
    """Approve a format-verified document and synchronize linked submission."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    UPDATE documents
    SET approval_status = 'Approved',
        approved_by = ?,
        approved_at = CURRENT_TIMESTAMP
    WHERE id = ?
    ''', (approved_by_user_id, doc_id))
    
    # Sync submission status if linked
    cursor.execute("SELECT submission_id FROM documents WHERE id = ?", (doc_id,))
    row = cursor.fetchone()
    if row and row['submission_id']:
        cursor.execute("UPDATE submissions SET status = 'Evaluated' WHERE id = ?", (row['submission_id'],))
        
    conn.commit()
    conn.close()
    return True


def reject_document(doc_id: int, reason: str = "") -> bool:
    """Mark document as needing revisions or rejected."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    UPDATE documents
    SET approval_status = 'Needs Revision',
        executive_verdict = ?
    WHERE id = ?
    ''', (f"Revision Requested: {reason}" if reason else "Revision requested by mentor.", doc_id))
    
    # Sync submission status if linked
    cursor.execute("SELECT submission_id FROM documents WHERE id = ?", (doc_id,))
    row = cursor.fetchone()
    if row and row['submission_id']:
        cursor.execute("UPDATE submissions SET status = 'Needs Revision' WHERE id = ?", (row['submission_id'],))
        
    conn.commit()
    conn.close()
    return True


def seed_sample_documents_if_empty():
    """Initializes sample intern submissions and comparative inspection benchmark if empty."""
    init_documents_table()
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute("SELECT COUNT(*) as count FROM documents WHERE format_comparison_matrix_json IS NOT NULL AND format_comparison_matrix_json != '[]'")
    row = cursor.fetchone()
    if row and row['count'] > 0:
        conn.close()
        return

    # Check if we have an intern user and assignment
    cursor.execute("SELECT u.id, u.full_name FROM users u WHERE u.role = 'intern' LIMIT 1")
    intern_row = cursor.fetchone()
    
    cursor.execute("SELECT a.id, a.project_id, p.title as project_title FROM assignments a JOIN projects p ON a.project_id = p.id LIMIT 1")
    assign_row = cursor.fetchone()
    
    intern_id = intern_row['id'] if intern_row else 2
    intern_name = intern_row['full_name'] if intern_row else "Bhumika"
    assignment_id = assign_row['id'] if assign_row else 1
    project_id = assign_row['project_id'] if assign_row else 1
    project_title = assign_row['project_title'] if assign_row else "TrendLens AI Platform"
    
    # Create sample submission if submissions table is empty
    cursor.execute("SELECT id FROM submissions LIMIT 1")
    if not cursor.fetchone():
        cursor.execute('''
        INSERT INTO submissions (assignment_id, intern_id, title, file_path, file_type, notes, status)
        VALUES (?, ?, 'CaptionAI Milestone Deliverable', 'uploads/documents/CaptionAI.docx', 'docx', 'Completed Milestone 1 baseline and initial model setup.', 'Needs Revision')
        ''', (assignment_id, intern_id))
        sub_id = cursor.lastrowid
    else:
        cursor.execute("SELECT id FROM submissions LIMIT 1")
        sub_id = cursor.fetchone()['id']

    # Seed initial comparison record (Matching the CaptionAI milestone report analysis)
    missing_secs = [
        "2. System Architecture & Component Diagram",
        "5. Implementation Details for Milestones 2 through 6",
        "6. Testing Suite & Latency Benchmarks"
    ]
    
    sec_comps = [
        {
            "sample_section": "1. Executive Summary & Abstract",
            "intern_status": "Partially Present",
            "sample_expectation": "Formal structured abstract highlighting problem scope, solution architecture, and key deliverables.",
            "intern_coverage": "Contains short project description paragraph but lacks formal problem scope and executive summary headers.",
            "missing_details": "Add structured Executive Summary highlighting key deliverables."
        },
        {
            "sample_section": "2. System Architecture & Diagram",
            "intern_status": "Missing",
            "sample_expectation": "Modular system architecture diagram, pipeline mapping, and technology stack breakdown.",
            "intern_coverage": "The architectural description is missing and the document abruptly cuts off under Activity 1.3.",
            "missing_details": "Draw and document full system architecture diagram and component interaction flow."
        },
        {
            "sample_section": "3. Implementation Milestones (1 to 6)",
            "intern_status": "Partially Present",
            "sample_expectation": "Step-by-step implementation logs and code outcomes for all 6 project milestones.",
            "intern_coverage": "Contains partial notes for Milestone 1; Milestones 2-6 are completely missing despite being listed in the workflow outline.",
            "missing_details": "Provide detailed implementation writeups for Milestones 2 through 6."
        },
        {
            "sample_section": "4. Testing, Metrics & Validation",
            "intern_status": "Missing",
            "sample_expectation": "Unit test results, accuracy benchmarks, and latency evaluation tables.",
            "intern_coverage": "No testing results or validation metrics provided.",
            "missing_details": "Add test execution table with pass/fail ratios."
        }
    ]
    
    fmt_checks = [
        {"check_name": "Title & Executive Summary", "passed": False, "details": "The document contains a 'Project Description' but lacks a formal, structured 'Executive Summary' or 'Abstract' section."},
        {"check_name": "Numbered Section Hierarchy", "passed": False, "details": "The document uses inconsistent and unnumbered primary headers ('Scenarios', 'Technical Architecture', 'Project Workflow')."},
        {"check_name": "Methodology & Architecture Depth", "passed": False, "details": "The architectural description is missing and the document abruptly cuts off under Activity 1.3 ('Draft an Architectural Diagram')."},
        {"check_name": "Implementation & Results", "passed": False, "details": "Activities 2 through 6 are completely missing from the actual content due to truncation, despite being listed in the workflow outline."},
        {"check_name": "References & Citations", "passed": True, "details": "Standard citations and tool references provided."},
        {"check_name": "Typography & Formatting Cleanliness", "passed": True, "details": "Clean typography without corrupted characters."}
    ]
    
    feedbacks = [
        "Step 1: Add the missing 'System Architecture & Diagram' section with a clean modular diagram as demonstrated in the mentor sample document.",
        "Step 2: Complete Implementation Milestones 2 through 6 instead of ending at Activity 1.3.",
        "Step 3: Include a formal structured Executive Summary at the start of the document.",
        "Step 4: Adopt consistent numbered section hierarchy (1.0, 2.0, 3.0) matching enterprise standards."
    ]

    format_matrix = [
        {
            "dimension": "Document Header & Title",
            "sample_standard": "Formal Enterprise Title with Version, Author, Track, and Date metadata block",
            "intern_format": "Informal heading 'CaptionAI' with no version, date, or author metadata block",
            "status": "Non-Compliant",
            "recommendation": "Add structured document metadata header with version, author, track, and date."
        },
        {
            "dimension": "Numbered Section Hierarchy",
            "sample_standard": "Strict 2-level numbering (1.0, 1.1, 2.0, 2.1, 3.0, 4.0, 5.0)",
            "intern_format": "Unnumbered text headings ('Scenarios', 'Technical Architecture', 'Project Workflow')",
            "status": "Non-Compliant",
            "recommendation": "Reformat all section titles using strict decimal numbering (1.0, 1.1, 2.0)."
        },
        {
            "dimension": "Executive Summary & Abstract",
            "sample_standard": "Formal Executive Summary outlining problem statement and business value",
            "intern_format": "Has informal 'Project Description' but lacks structured abstract and business need",
            "status": "Non-Compliant",
            "recommendation": "Include formal 1.0 Executive Summary & 1.1 Problem Statement."
        },
        {
            "dimension": "Architecture & Flow Diagrams",
            "sample_standard": "Modular ASCII/Mermaid or image architecture diagram with component specifications table",
            "intern_format": "Cuts off abruptly under Activity 1.3 'Draft an Architectural Diagram' without any diagram",
            "status": "Missing",
            "recommendation": "Include complete high-level system architecture diagram and component table."
        },
        {
            "dimension": "Implementation Milestones (1 to 6)",
            "sample_standard": "Full tabular breakdown and implementation notes for all 6 milestones",
            "intern_format": "Only contains partial notes for Milestone 1; Milestones 2-6 completely missing",
            "status": "Missing",
            "recommendation": "Provide completed progress writeups and status indicators for Milestones 2 through 6."
        },
        {
            "dimension": "Testing & Citations",
            "sample_standard": "Automated test matrix table and formal IEEE-style reference list",
            "intern_format": "Contains basic links but missing structured test verification matrix table",
            "status": "Partially Compliant",
            "recommendation": "Add Testing & Validation Matrix table with pass criteria and metrics."
        }
    ]

    proper_template = """# CaptionAI — Automated Multimodal Video & Image Captioning System
**Document Type**: Milestone Deliverable & Technical Specification | **Version**: 1.0
**Author**: John Intern | **Track**: AI & Machine Learning | **Date**: September 2026

---

## 1.0 Executive Summary & Project Abstract
### 1.1 Problem Statement & Industry Need
In digital media workflows, content creators struggle with manual accessibility compliance and caption generation. CaptionAI provides an automated, multimodal video and image captioning pipeline powered by transformer models.

### 1.2 Core Solution Architecture
CaptionAI integrates a FastAPI backend with PyTorch BLIP-2 vision-language models, streaming real-time generated subtitles and structured metadata.

---

## 2.0 System Architecture & Component Design
### 2.1 High-Level Architecture Diagram
```
[User Interface / Video Ingestion]
                │ (HTTP Multipart)
                ▼
[FastAPI Ingestion & Audio-Visual Splitter]
                │
    ┌───────────┴───────────┐
    ▼                       ▼
[Whisper Audio ASR]   [BLIP-2 Visual Frame Extractor]
    └───────────┬───────────┘
                ▼
  [Multimodal Caption Fusion Engine]
                │
                ▼
 [Subtitle Exporter (.SRT / .VTT / JSON)]
```

### 2.2 Component Specifications
| Component | Technology | Responsibility |
|---|---|---|
| Ingestion & Preprocessing | FFmpeg & OpenCV | Frame extraction & audio track splitting |
| Visual Captioning Core | Salesforce BLIP-2 | Deep visual understanding & scene captioning |
| Audio Transcription | OpenAI Whisper | High-accuracy speech-to-text alignment |
| Synchronization Engine | Python NumPy | Timestamp reconciliation & subtitle generation |

---

## 3.0 Implementation Milestones & Work Breakdown
| Milestone | Objective | Deliverables | Verification Status |
|---|---|---|---|
| Milestone 1.0 | Environment & Dataset Setup | Requirements.txt, Model weights cache | ✅ Completed |
| Milestone 2.0 | Video Preprocessing Pipeline | `preprocess.py`, FFmpeg pipeline | ⏳ In Progress |
| Milestone 3.0 | Multimodal Inference Engine | `model_pipeline.py`, BLIP-2 integration | ⏳ In Progress |
| Milestone 4.0 | REST API & WebSocket Stream | FastAPI server, Async processing | ⏳ Pending |
| Milestone 5.0 | Testing & Latency Benchmarks | Unit test suite, 95% pass rate | ⏳ Pending |
| Milestone 6.0 | Deployment & Containerization | Dockerfile, CI/CD pipeline | ⏳ Pending |

---

## 4.0 Testing, Metrics & Validation Suite
### 4.1 Automated Test Execution Matrix
| Test Suite | Target Module | Pass Criteria | Result |
|---|---|---|---|
| Frame Extraction Test | `preprocess.py` | Extract 10 FPS without dropped frames | PASS |
| Model Latency Test | `inference.py` | Inference latency < 250ms per keyframe | PASS |
| Subtitle Alignment | `sync.py` | BLEU / CIDEr alignment score > 0.82 | PASS |

---

## 5.0 References & Enterprise Standards
[1] Li et al., "BLIP-2: Bootstrapping Language-Image Pre-training with Frozen LLMs", 2023.
[2] Radford et al., "Robust Speech Recognition via Large-Scale Weak Supervision (Whisper)", 2022.
[3] IEEE Technical Documentation Standard for AI/ML Deliverables (IEEE 26514-2022)."""

    cursor.execute('''
    INSERT INTO documents (
        title, file_path, file_name, file_type, file_size_kb, category,
        project_id, uploaded_by, submission_id, intern_id, intern_name,
        sample_file_path, sample_file_name, extracted_text_preview,
        extracted_sections_json, formatting_checks_json, missing_sections_json,
        comparison_details_json, format_comparison_matrix_json, proper_format_template,
        intern_actionable_feedback, format_score, format_status, approval_status,
        executive_verdict
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        "CaptionAI",
        "uploads/documents/CaptionAI.docx",
        "CaptionAI.docx",
        "docx",
        3284.4,
        "Milestone Report",
        project_id,
        1,
        sub_id,
        intern_id,
        intern_name,
        "uploads/samples/Sample_Milestone_Report_Standard.docx",
        "Sample_Milestone_Report_Standard.docx",
        "CaptionAI Project Deliverable Document Preview...",
        json.dumps(sec_comps),
        json.dumps(fmt_checks),
        json.dumps(missing_secs),
        json.dumps(sec_comps),
        json.dumps(format_matrix),
        proper_template,
        json.dumps(feedbacks),
        45,
        "Missing Required Sections",
        "Pending Review",
        "The document cannot be approved as it is severely truncated mid-sentence, lacks a completed system architecture section, contains no implementation details for Milestones 2–6, and does not follow standard enterprise numbering conventions. Please request a complete and properly formatted document from the student."
    ))
    
    conn.commit()
    conn.close()
